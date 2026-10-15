from __future__ import annotations

import uuid
from io import BytesIO
from types import SimpleNamespace

from docx import Document

from mvb.export.docx import build_summary_docx
from mvb.schemas.analysis import AnalysisBundleOut


def _bundle(**overrides) -> AnalysisBundleOut:
    base = {
        "meeting_id": uuid.uuid4(),
        "summary": {"text": "The team reviewed the launch plan.", "bullets": ["Ship on Friday", "Budget approved"]},
    }
    base.update(overrides)
    return AnalysisBundleOut(**base)


def test_docx_structure_and_styles() -> None:
    meeting = SimpleNamespace(title="Launch Planning", started_at=None, duration_ms=3_600_000)
    data = build_summary_docx(meeting=meeting, bundle=_bundle())
    doc = Document(BytesIO(data))
    by_style = {p.style.name: [] for p in doc.paragraphs}
    for p in doc.paragraphs:
        by_style.setdefault(p.style.name, []).append(p.text)
    assert "Launch Planning" in by_style["Title"]
    assert by_style["Heading 1"] == ["Summary"]
    assert "The team reviewed the launch plan." in [p.text for p in doc.paragraphs]
    assert by_style["List Bullet"] == ["Ship on Friday", "Budget approved"]


def test_docx_labelled_mock_note() -> None:
    data = build_summary_docx(
        meeting=SimpleNamespace(title="M", started_at=None, duration_ms=0), bundle=_bundle(mock=True)
    )
    doc = Document(BytesIO(data))
    assert any("mock" in p.text.lower() for p in doc.paragraphs)


def test_docx_omits_absent_sections() -> None:
    data = build_summary_docx(meeting=SimpleNamespace(title="M", started_at=None, duration_ms=0), bundle=_bundle())
    doc = Document(BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Decisions" not in headings and "Action items" not in headings and "Q&A" not in headings
