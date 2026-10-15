from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document

from mvb.schemas.analysis import AnalysisBundleOut


def _fmt_ms(ms: int) -> str:
    seconds = max(0, int(ms)) // 1000
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def _items(payload: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not isinstance(payload, dict):
        return []
    raw = payload.get("items")
    if not isinstance(raw, list):
        return []
    return [item for item in raw if isinstance(item, dict)]


def build_summary_docx(*, meeting: Any, bundle: AnalysisBundleOut) -> bytes:
    """Render the analysis bundle as a structured .docx.

    Rules: sections with no data are OMITTED (never invented); mock output is labelled.
    """
    doc = Document()
    doc.add_heading(meeting.title or "Meeting", level=0)

    meta: list[str] = []
    if getattr(meeting, "started_at", None):
        meta.append(meeting.started_at.strftime("%Y-%m-%d %H:%M"))
    if getattr(meeting, "duration_ms", 0):
        meta.append(f"{_fmt_ms(meeting.duration_ms)} long")
    doc.add_paragraph(" · ".join(meta) if meta else "Meeting summary")

    if bundle.mock:
        note = doc.add_paragraph().add_run(
            "Note: this analysis was produced by the labelled mock provider - the content is "
            "simulated, not real meeting analysis."
        )
        note.italic = True

    summary = bundle.summary if isinstance(bundle.summary, dict) else None
    if summary:
        doc.add_heading("Summary", level=1)
        text = str(summary.get("text") or "").strip()
        if text:
            doc.add_paragraph(text)
        for bullet in summary.get("bullets") or []:
            doc.add_paragraph(str(bullet), style="List Bullet")

    if bundle.sections:
        doc.add_heading("Agenda & discussion", level=1)
        section_docs = {d.section_id: d for d in bundle.section_summaries}
        for section in bundle.sections:
            heading = section.title
            if section.start_ms or section.end_ms:
                heading += f"  [{_fmt_ms(section.start_ms)}-{_fmt_ms(section.end_ms)}]"
            doc.add_heading(heading, level=2)
            if section.summary:
                doc.add_paragraph(section.summary)
            section_doc = section_docs.get(section.id)
            payload = section_doc.payload if section_doc and isinstance(section_doc.payload, dict) else {}
            for key in ("bullets", "decisions", "action_items"):
                for entry in payload.get(key) or []:
                    doc.add_paragraph(str(entry), style="List Bullet")

    decisions = _items(bundle.decisions)
    if decisions:
        doc.add_heading("Decisions", level=1)
        for item in decisions:
            text = str(item.get("text") or "").strip()
            if not text:
                continue
            line = text
            by = item.get("decided_by")
            if isinstance(by, list) and by:
                line += " - " + ", ".join(str(x) for x in by)
            at_ms = item.get("at_ms")
            if isinstance(at_ms, (int, float)):
                line += f"  [{_fmt_ms(int(at_ms))}]"
            doc.add_paragraph(line, style="List Bullet")

    actions = _items(bundle.action_items)
    if actions:
        doc.add_heading("Action items", level=1)
        for item in actions:
            text = str(item.get("text") or "").strip()
            if not text:
                continue
            line = text
            if item.get("owner"):
                line += f" - owner: {item['owner']}"
            if item.get("due"):
                line += f" - due: {item['due']}"
            doc.add_paragraph(line, style="List Bullet")

    if bundle.qa:
        doc.add_heading("Q&A", level=1)
        for qa in bundle.qa:
            payload = qa.payload if isinstance(qa.payload, dict) else {}
            question = str(payload.get("question") or "").strip()
            answer = str(payload.get("answer") or "").strip()
            if not question:
                continue
            paragraph = doc.add_paragraph()
            paragraph.add_run(question).bold = True
            if answer:
                doc.add_paragraph(answer)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
